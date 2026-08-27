"""
Validation: skincarisma and amazon, side by side, every product

Why this exists
  My supervisors said the work lacks validation, and that a summary statistic
  is not enough. They want to see the two sources placed next to each other and
  compared product by product, not a sample.

Usage:
    py build_validation_side_by_side.py
"""
import re
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SC_FILE  = 'skincarisma_ALL_products.csv'
DS_FILE  = 'OPTION_A_mixed_sources.csv'
AMZ_FILE = 'amazon_skintype.csv'
MAP_FILE = 'acc19k_matches.csv'
OUT      = 'VALIDATION_SkinCarisma_vs_Amazon.xlsx'

norm = lambda s: re.sub(r'[^a-z0-9]', '', str(s).lower())

# ---------------------------------------------------------------- load
sc = pd.read_csv(SC_FILE, low_memory=False, dtype=str).fillna('')
sc = sc[sc['found'] == '1'].drop_duplicates('product_id')
ds = pd.read_csv(DS_FILE, low_memory=False, dtype=str).fillna('')
amz = pd.read_csv(AMZ_FILE, dtype=str).fillna('')
mp = pd.read_csv(MAP_FILE, low_memory=False, dtype=str).fillna('')
print(f'SkinCarisma products with a reading : {len(sc):,}')

# attach the Amazon ASIN and the raw declared text to every product
mp['_k'] = mp['brand'].map(norm) + '|' + mp['name'].map(norm)
ds['_k'] = ds['brand'].map(norm) + '|' + ds['name'].map(norm)
ds['amazon_asin'] = ds['_k'].map(dict(zip(mp['_k'], mp['am_asin'])))
ds['amazon_raw'] = ds['amazon_asin'].map(dict(zip(amz['asin'], amz['amazon_skin_type_raw'])))

# NOTE: sc and ds BOTH have brand and name, so merging without care produces
# brand_x / brand_y and the readable columns vanish from the output. Take the
# identifying fields from the dataset side only, renamed up front.
keep = ['product_id', 'brand', 'name', 'product_type', 'amazon_asin', 'amazon_raw']
right = ds[[c for c in keep if c in ds.columns]].copy()
df = sc.drop(columns=[c for c in ('brand', 'name', 'product_type') if c in sc.columns]) \
       .merge(right, on='product_id', how='left').fillna('')
assert 'brand' in df.columns and 'name' in df.columns, 'brand/name lost in the merge'
print(f'of those, with an Amazon declaration   : {int((df["amazon_raw"] != "").sum()):,}')


# ---------------------------------------------------------------- parse
def amz_axis(raw, axis):
    """yes / no / '' for one axis, from the declared Amazon text."""
    s = str(raw).strip().lower()
    if s == '':
        return ''
    if s in ('all', 'all skin types', 'all skin type', 'universal', 'any'):
        return ''                       # a universal claim makes no axis claim
    if axis == 'dry':
        return 'YES' if 'dry' in s else 'no'
    if axis == 'oily':
        return 'YES' if ('oil' in s or 'acne' in s) else 'no'
    return 'YES' if 'sensitiv' in s else 'no'


def sc_axis(v):
    return {'1': 'YES', '0': 'no'}.get(str(v).strip(), '')


def verdict(a, b):
    if a == '' or b == '':
        return 'not comparable'
    return 'AGREE' if a == b else 'DISAGREE'


UNIVERSAL = re.compile(r'^\s*(all|all skin types?|universal|any)\s*$', re.I)
df['amazon_is_universal'] = df['amazon_raw'].str.match(UNIVERSAL).map({True: 'YES', False: ''}).fillna('')

for axis, col in (('dry', 'sc_dry'), ('oily', 'sc_oily'), ('sensitive', 'sc_sensitive')):
    df[f'SC_{axis}'] = df[col].map(sc_axis)
    df[f'AMZ_{axis}'] = df['amazon_raw'].map(lambda r, a=axis: amz_axis(r, a))
    df[f'{axis}_verdict'] = [verdict(x, y) for x, y in zip(df[f'SC_{axis}'], df[f'AMZ_{axis}'])]

vcols = ['dry_verdict', 'oily_verdict', 'sensitive_verdict']
df['axes_compared'] = (df[vcols] != 'not comparable').sum(axis=1)
df['axes_agreeing'] = (df[vcols] == 'AGREE').sum(axis=1)


def overall(r):
    if r['amazon_raw'] == '':
        return '- no Amazon record'
    if r['amazon_is_universal'] == 'YES':
        return '- Amazon says "All" (no claim)'
    if r['axes_compared'] == 0:
        return '- nothing comparable'
    if r['axes_agreeing'] == r['axes_compared']:
        return 'FULL AGREEMENT'
    if r['axes_agreeing'] == 0:
        return 'FULL DISAGREEMENT'
    return f"PARTIAL {r['axes_agreeing']}/{r['axes_compared']}"


df['OVERALL'] = df.apply(overall, axis=1)

# a readable single label from each side, so a human can scan the sheet
def sc_label(r):
    d, o, s = r['SC_dry'] == 'YES', r['SC_oily'] == 'YES', r['SC_sensitive'] == 'YES'
    t = 'Combination' if (d and o) else 'Oily' if o else 'Dry' if d else 'Normal'
    return t + (' + Sensitive' if s else '')


df['SkinCarisma_says'] = df.apply(sc_label, axis=1)
df['Amazon_says'] = df['amazon_raw']

# ---------------------------------------------------------------- sheets
FULL = [
    'product_id', 'brand', 'name', 'product_type',
    'SkinCarisma_says', 'Amazon_says', 'OVERALL',
    'SC_dry', 'AMZ_dry', 'dry_verdict',
    'SC_oily', 'AMZ_oily', 'oily_verdict',
    'SC_sensitive', 'AMZ_sensitive', 'sensitive_verdict',
    'sc_dry_good', 'sc_dry_bad', 'sc_oily_good', 'sc_oily_bad',
    'sc_sensitive_good', 'sc_sensitive_bad',
    'amazon_asin', 'amazon_is_universal', 'axes_compared', 'axes_agreeing', 'sc_url',
]
full = df[[c for c in FULL if c in df.columns]].copy()

# ---- summary ---------------------------------------------------------------
n_sc = len(df)
n_amz = int((df['amazon_raw'] != '').sum())
n_uni = int((df['amazon_is_universal'] == 'YES').sum())
n_cmp = int((df['axes_compared'] > 0).sum())

rows = [['POPULATION', '', ''],
        ['Products with a SkinCarisma reading', n_sc, '100.0%'],
        ['  of those, with an Amazon declaration', n_amz, f'{100*n_amz/n_sc:.1f}%'],
        ['  of those, Amazon declares only "All"', n_uni, f'{100*n_uni/n_sc:.1f}%'],
        ['  COMPARABLE on at least one axis', n_cmp, f'{100*n_cmp/n_sc:.1f}%'],
        ['  no Amazon record at all', n_sc - n_amz, f'{100*(n_sc-n_amz)/n_sc:.1f}%'],
        ['', '', '']]

rows.append(['AGREEMENT PER AXIS  (only rows where both sources speak)', '', ''])
for axis in ('dry', 'oily', 'sensitive'):
    v = df[f'{axis}_verdict']
    cmpable = df[v != 'not comparable']            # count ONLY where both sources speak
    c = len(cmpable)
    a = int((v == 'AGREE').sum())
    both_yes = int(((cmpable[f'SC_{axis}'] == 'YES') & (cmpable[f'AMZ_{axis}'] == 'YES')).sum())
    amz_yes = int((cmpable[f'AMZ_{axis}'] == 'YES').sum())
    sc_yes = int((cmpable[f'SC_{axis}'] == 'YES').sum())
    rows.append([f'{axis}: compared / agree / agreement rate', f'{c:,} / {a:,}',
                 f'{100*a/c:.1f}%' if c else 'n/a'])
    rows.append([f'   {axis}: Amazon YES / SkinCarisma YES / both YES',
                 f'{amz_yes:,} / {sc_yes:,} / {both_yes:,}',
                 f'recall {100*both_yes/amz_yes:.1f}%' if amz_yes else 'n/a'])
rows.append(['', '', ''])

rows.append(['PER-PRODUCT VERDICT', '', ''])
for k, v in df['OVERALL'].value_counts().items():
    rows.append([f'  {k}', v, f'{100*v/n_sc:.1f}%'])

summary = pd.DataFrame(rows, columns=['Measure', 'Value', 'Share'])

dis = full[full['OVERALL'].str.contains('DISAGREE|PARTIAL', na=False)]

method = pd.DataFrame([
    ['What this workbook is',
     'Every one of the 6,627 products SkinCarisma returned a reading for, placed next to the Amazon declaration for the same product, compared axis by axis. Not a sample.'],
    ['Why every product is shown',
     'Products with no Amazon counterpart are kept and marked, so the reader sees the whole population rather than only the comparable part.'],
    ['How SkinCarisma is read',
     'It counts good vs bad ingredients per axis. "Dry Skin 6 good / 2 bad" gives 6 > 2, so dry = YES. The raw counts are in the sheet so every call can be checked.'],
    ['How Amazon is read',
     'The declared details["Skin Type"] text is parsed. "Dry" gives dry = YES. "Oily, Combination" gives oily = YES.'],
    ['Why "All" is excluded from the rate',
     'A product declared "All" agrees with every possible value by construction, so counting it would inflate every figure. Following Vendruscolo et al. (2025), Dermatological Reviews 6:e70045, "all skin types" is a tolerance-testing claim, not a classification. Those rows are shown in full and flagged, but not scored.'],
    ['What "not comparable" means',
     'One of the two sources has no reading on that axis, so no comparison is possible. No value is invented to fill it.'],
    ['What disagreement means',
     'Not that one source is broken. A manufacturer writing "for sensitive skin" describes who they sell it to. SkinCarisma counts irritant-flagged ingredients in the formula. Two different questions.'],
], columns=['Point', 'Explanation'])

# ---------------------------------------------------------------- write
GREEN, RED, AMBER = 'E4F0E4', 'FBE9E9', 'FDF3E6'
with pd.ExcelWriter(OUT, engine='openpyxl') as w:
    summary.to_excel(w, sheet_name='1 Summary', index=False)
    full.to_excel(w, sheet_name='2 Full comparison', index=False)
    dis.to_excel(w, sheet_name='3 Disagreements only', index=False)
    method.to_excel(w, sheet_name='4 Method', index=False)

    for nm in w.sheets:
        ws = w.sheets[nm]
        for c in ws[1]:
            c.fill = PatternFill('solid', fgColor='584A7A')
            c.font = Font(bold=True, color='FFFFFF', size=10)
            c.alignment = Alignment(vertical='center', wrap_text=True)
        ws.row_dimensions[1].height = 30
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions

    # colour the verdict cells so a reader can scan a 6,627 row sheet
    ws = w.sheets['2 Full comparison']
    cols = {c.value: c.column for c in ws[1]}
    for name in ('dry_verdict', 'oily_verdict', 'sensitive_verdict', 'OVERALL'):
        if name not in cols:
            continue
        L = get_column_letter(cols[name])
        for i in range(2, ws.max_row + 1):
            cell = ws[f'{L}{i}']
            v = str(cell.value or '')
            if 'DISAGREE' in v:
                cell.fill = PatternFill('solid', fgColor=RED)
            elif 'AGREE' in v:
                cell.fill = PatternFill('solid', fgColor=GREEN)
            elif 'PARTIAL' in v:
                cell.fill = PatternFill('solid', fgColor=AMBER)

    widths = {'brand': 20, 'name': 44, 'product_type': 18, 'SkinCarisma_says': 20,
              'Amazon_says': 30, 'OVERALL': 26, 'sc_url': 40, 'Measure': 52,
              'Value': 20, 'Share': 12, 'Point': 30, 'Explanation': 95}
    for nm in w.sheets:
        ws = w.sheets[nm]
        for c in ws[1]:
            ws.column_dimensions[get_column_letter(c.column)].width = widths.get(str(c.value), 13)

# ---------------------------------------------------------------- report
print('\n' + '=' * 70)
print(summary.to_string(index=False))
print('=' * 70)
print(f'\nwrote {OUT}')
print(f'  sheet 2 holds all {len(full):,} products')
print(f'  sheet 3 holds the {len(dis):,} rows where the sources conflict')
