"""
The lebanese origin dataset: brands that are made in lebanon

What went wrong in the harvest, and how it is fixed
  38,835 products were harvested from 82 sites. Most of them do not belong
  here, because the discovery step found Lebanese WEBSITES rather than Lebanese
  BRANDS. It checked that a site was Lebanese and mentioned skincare words. It
  never checked that the site belongs to a company that MAKES something.

  So the file contains:

      abedtahan.com      wine coolers, Lenovo laptops, Atari consoles
      pcandparts.com     "#1 Computer Shop in Lebanon", 53% computer parts
      lebzone.com        books, gift ideas, 16% food
      sadaf.com          Middle Eastern groceries, 0% skincare
      fattalonline.com   a perfume distributor carrying 136 brands

Usage:
    py build_lebanese_origin.py
"""
import re
import unicodedata
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HARVEST = 'lebanese_origin_harvested.csv'
OLD = 'ALL_SOURCES.xlsx'          # the original 332 rows
OUT_C = 'LEBANESE_ORIGIN.csv'
OUT_X = 'LEBANESE_ORIGIN.xlsx'
DROPPED = 'LEBANESE_ORIGIN_dropped.csv'

SCHEMA = ['product_id', 'brand', 'name', 'product_type', 'country',
          'skin_type', 'sensitivity', 'skin_type_status', 'skin_type_tier',
          'skin_type_authority', 'skin_type_source', 'skin_type_rule',
          'skin_type_quote', 'skin_type_url', 'ingredients', 'ingredient_count',
          'key_ingredients', 'free_from', 'spf', 'benefits', 'concerns',
          'rating', 'review_count', 'review_source', 'review_texts_json',
          'product_summary', 'skinsort_url']
EXTRA = ['brand_country', 'market', 'price_usd', 'product_url', 'domain',
         'harvest_method', 'origin_evidence']

# ---------------------------------------------------- hand-checked decisions
# Confirmed Lebanese makers: the company designs and produces the product.
KEEP_DOMAINS = {
    'beesline.com', 'khanelkaser.com', 'lb.khanelkaser.com', 'thealoelab.com',
    'helwe.com', 'samasoaps.com', 'senteursdorient.com', 'savonduliban.com',
    'casia-handmade.com', 'cherryblossomleb.com', 'atelierbeautanique.com',
    'nurayacosmetics.com', 'paradise-cosmetic.com', 'splashycosmetics.com',
    'masbanatawaida.com', 'lb.houseofsoap.com', 'jana-lb.com',
    'laterratales.com', 'darmmess.com', 'tropcosmetics.com',
    'sbrandsofficial.com', 'xiranskincare.com', 'cosmaline.com',
    'shop.cosmaline.com', 'zejd.net', 'duft-lb.com', 'ecladerm.com',
    'aloelab.com', 'yvesmorel.com',
}
# Not Lebanese makers, with the reason recorded so the choice is defensible.
DROP_DOMAINS = {
    'abedtahan.com': 'an appliance and electronics shop',
    'pcandparts.com': 'a computer shop',
    'lebzone.com': 'books and gifts',
    'sadaf.com': 'a Middle Eastern grocery',
    'livgood.com': 'a food and wellness shop',
    'healthy961.com': 'a health food shop',
    'terroirsduliban.com': 'Lebanese food producers',
    'lecomptoirduliban.com': 'a Lebanese grocery in Paris',
    'saveursliban-lecomptoir.com': 'a Lebanese grocery',
    'lebanesesignature.com': 'Lebanese food and gifts',
    'loubnany.com': 'Lebanese food',
    'made-in-lebanon.net': 'a manufacturers directory, not a brand',
    'personalcarecouncil.org': 'a trade association',
    'berytech.org': 'a technology incubator',
    'euromonitor.com': 'a market research firm',
    'statista.com': 'a statistics site',
    'fr.ankorstore.com': 'a wholesale marketplace',
    'boutique-artisans-du-monde.com': 'a fair trade shop, not Lebanese',
    'beirutboutique.co.uk': 'a gift shop',
    'lushlebanon.com': 'Lush is a British brand, this is Lebanese retail',
    'flormarlebanon.com': 'Flormar is a Turkish brand, this is Lebanese retail',
    'thebodyshop.com.lb': 'The Body Shop is British, this is Lebanese retail',
    'xirancosmetics.com': 'a contract manufacturer offering to make other brands',
    'juicybeauty.online': 'a retailer',
    'gomarbeauty.com': 'a retailer, and based in the UAE',
    'parapharm.com.lb': 'a pharmacy retailer',
    'beautedulibanmarche.com': 'a marketplace carrying many brands',
    'waw.beauty': 'a retailer',
    'faceshine.me': 'a retailer',
    'hbytala.com': 'a retailer',
    'shop.cosmaline.com': '',      # handled by KEEP, kept for clarity
}
DROP_DOMAINS.pop('shop.cosmaline.com', None)

# page titles that ended up in the vendor field when a site had none
TITLE_JUNK = re.compile(
    r'^(home|homepage|bienvenue|welcome|one|shop|store|boutique|accueil|'
    r'products?|collections?|all products|untitled|index|main|'
    r'#?\d*\s*(computer|shop)\b.*)$', re.I)

SKIN = re.compile(
    r'\b(serum|cream|creme|moisturi[sz]\w*|cleanser|toner|mask|scrub|exfoliat\w*|'
    r'sunscreen|spf|balm|lotion|face|facial|skin|soap|savon|gel|essence|ampoule|'
    r'peeling|micellar|hydrat\w*|anti[- ]?ag\w+|acne|blemish|butter|body oil|'
    r'shower|deodorant|hand cream|foot cream|lip)\b', re.I)
NOTSKIN = re.compile(
    r'\b(zaatar|honey|jam|molasses|tahini|halva|coffee|tea\b|wine|arak|olives|'
    r'spice|sumac|freekeh|bulgur|rice|nuts|chocolate|biscuit|syrup|vinegar|'
    r'labneh|cheese|dates|maamoul|laptop|lenovo|ssd|monitor|printer|\btv\b|'
    r'fridge|cooler|gaming|console|phone|tablet|camera|headphone|keyboard|'
    r'book|candle|mug|shirt|bag\b|jewel|earring|necklace|toy|furniture|'
    r'gift card|voucher|perfume|eau de parfum|eau de toilette|fragrance mist)\b',
    re.I)

CATEGORY = [
    (re.compile(r'serum|ampoule|concentrate', re.I), 'Serum'),
    (re.compile(r'cleanser|face wash|foaming|micellar|cleansing gel', re.I), 'Face Cleanser'),
    (re.compile(r'sunscreen|spf|sun cream|sun block', re.I), 'Sunscreen'),
    (re.compile(r'toner|tonic|rose water|floral water', re.I), 'Toner'),
    (re.compile(r'sheet mask', re.I), 'Sheet Mask'),
    (re.compile(r'mask|masque', re.I), 'Wet Mask'),
    (re.compile(r'scrub|exfoliat|peel|gommage', re.I), 'Exfoliator'),
    (re.compile(r'eye cream|eye serum|eye contour|eye gel', re.I), 'Eye Moisturizer'),
    (re.compile(r'lip balm|lip butter|lip care|lip oil', re.I), 'Lip Moisturizer'),
    (re.compile(r'makeup remover|cleansing oil|cleansing balm', re.I), 'Makeup Remover'),
    (re.compile(r'\boil\b|huile', re.I), 'Oil'),
    (re.compile(r'essence', re.I), 'Essence'),
    (re.compile(r'emulsion', re.I), 'Emulsion'),
    (re.compile(r'hand cream|hand lotion|hand butter', re.I), 'Hand Care'),
    (re.compile(r'night cream|nuit', re.I), 'Night Moisturizer'),
    (re.compile(r'day cream|jour', re.I), 'Day Moisturizer'),
    (re.compile(r'soap|savon|bar\b', re.I), 'Bath & Body'),
    (re.compile(r'body|shower|bath', re.I), 'Bath & Body'),
    (re.compile(r'treatment|repair|corrector', re.I), 'Facial Treatment'),
    (re.compile(r'cream|creme|moistur|lotion|balm|butter', re.I), 'General Moisturizer'),
]


def asc(s):
    return unicodedata.normalize('NFKD', str(s)).encode('ascii', 'ignore').decode().lower()


def categorise(name, ptype):
    t = f'{name} {ptype}'
    for pat, cat in CATEGORY:
        if pat.search(t):
            return cat
    return ''


def brand_from_domain(dom):
    base = dom.split('.')[0]
    if base in ('lb', 'shop', 'www', 'fr', 'store'):
        base = dom.split('.')[1]
    base = re.sub(r'(cosmetics|skincare|beauty|lebanon|online|official)$', '', base)
    return base.replace('-', ' ').title().strip() or dom


d = pd.read_csv(HARVEST, dtype=str, low_memory=False).fillna('')
start = len(d)
print(f'{start:,} products harvested from {d["domain"].nunique()} sites\n')

# ------------------------------------------------- 1. is the site a maker?
vend = {}
for dom, s in d.groupby('domain'):
    vc = s['brand'].value_counts()
    vend[dom] = (len(s), s['brand'].nunique(),
                 round(100 * vc.iloc[0] / len(s)) if len(vc) else 0)

reasons = {}
keep_dom = set()
for dom, (n, nv, pct) in vend.items():
    if dom in DROP_DOMAINS:
        reasons[dom] = DROP_DOMAINS[dom]
    elif dom in KEEP_DOMAINS:
        keep_dom.add(dom)
    elif pct < 60:
        reasons[dom] = f'a shop, it carries {nv} different brands'
    elif n == 0:
        reasons[dom] = 'nothing harvested'
    else:
        # not on either list, and looks like a single-brand site
        sk = d[d['domain'] == dom]['name'].str.contains(SKIN, na=False).mean()
        if sk >= 0.30:
            keep_dom.add(dom)
        else:
            reasons[dom] = f'only {100*sk:.0f}% of its products are skincare'

print(f'sites kept as Lebanese makers   {len(keep_dom)}')
print(f'sites dropped                   {len(reasons)}\n')

d['keep_site'] = d['domain'].isin(keep_dom)

# ------------------------------------------- 2. is the PRODUCT skincare?
blob = d['name'] + ' ' + d['product_type']
d['is_skin'] = blob.str.contains(SKIN, na=False) & ~blob.str.contains(NOTSKIN, na=False)

keep = d[d['keep_site'] & d['is_skin']].copy()
drop = d[~(d['keep_site'] & d['is_skin'])].copy()
drop['why'] = [reasons.get(x, 'not a skincare product') for x in drop['domain']]
drop[['brand', 'name', 'domain', 'why']].to_csv(DROPPED, index=False)

print(f'products on a maker site        {int(d["keep_site"].sum()):,}')
print(f'of those, actually skincare     {len(keep):,}\n')

# --------------------------------------------------- 3. tidy up the columns
# A brand name is short and is not a product description. Sites with no vendor
# field fall back to the page title, which is sometimes a product:
#   "Serum to control Scalp Oiliness", "Face Balm for Very Dry Skin"
# Those would each become their own brand. Where the value looks like a product
# rather than a company, the brand is taken from the domain instead.
def looks_like_product(b):
    t = str(b).strip()
    if not t or TITLE_JUNK.match(t):
        return True
    if len(t) > 28 or len(t.split()) > 4:
        return True
    return bool(re.search(r'\b(serum|cream|balm|wash|lotion|mask|scrub|oil|'
                          r'cleanser|toner|spf|ml\b|shade|bundle|for)\b', t, re.I))


keep['brand'] = [brand_from_domain(dom) if looks_like_product(b) else str(b).strip()
                 for b, dom in zip(keep['brand'], keep['domain'])]
keep['product_type'] = [categorise(n, t) for n, t in
                        zip(keep['name'], keep['product_type'])]
keep = keep[keep['product_type'] != '']

for c in SCHEMA + EXTRA:
    if c not in keep.columns:
        keep[c] = ''
keep['country'] = 'Lebanon'
keep['brand_country'] = 'Lebanon'
keep['market'] = 'Lebanon'
keep['harvest_method'] = keep['method']
keep['origin_evidence'] = 'the brand runs this website and makes the product'
keep['ingredient_count'] = [str(len([x for x in str(i).split(',') if x.strip()]))
                            if str(i).strip() else '' for i in keep['ingredients']]
keep['skin_type_source'] = keep['domain']
keep['skin_type_url'] = keep['product_url']
keep.loc[keep['skin_type'] != '', 'skin_type_tier'] = '1'
keep.loc[keep['skin_type'] != '', 'skin_type_authority'] = 'manufacturer'
keep.loc[keep['skin_type'] != '', 'skin_type_rule'] = 'stated on the brand website'
keep.loc[keep['skin_type'] != '', 'skin_type_status'] = 'found'
keep.loc[keep['skin_type'] == '', 'skin_type_status'] = 'searched, not stated'

# ------------------------------------- 4. fold in the original 332 if present
try:
    x = pd.ExcelFile(OLD)
    sheet = [s for s in x.sheet_names if 'origin' in s.lower()][0]
    o = pd.read_excel(x, sheet_name=sheet, dtype=str).fillna('')
    o = o.rename(columns={'brand_name': 'brand', 'product_name': 'name',
                          'product_category': 'product_type',
                          'skin_concerns': 'concerns',
                          'product_url': 'product_url'})
    o['skin_type'] = o['skin_type'].replace({'Unknown': ''})
    o['country'] = 'Lebanon'
    o['brand_country'] = 'Lebanon'
    o['market'] = 'Lebanon'
    o['domain'] = [str(u).split('/')[2].replace('www.', '') if '://' in str(u) else ''
                   for u in o['product_url']]
    o['harvest_method'] = 'the original workbook'
    o['origin_evidence'] = 'listed as a Lebanese brand in the source workbook'
    o['product_type'] = [categorise(n, t) for n, t in zip(o['name'], o['product_type'])]
    o = o[o['product_type'] != '']
    for c in SCHEMA + EXTRA:
        if c not in o.columns:
            o[c] = ''
    both = pd.concat([keep[SCHEMA + EXTRA], o[SCHEMA + EXTRA]], ignore_index=True)
    print(f'the original workbook added     {len(o):,} rows')
except Exception as e:
    both = keep[SCHEMA + EXTRA].copy()
    print(f'(original workbook not merged: {e})')

# ------------------------------------------------------- 5. de-duplicate
SIZE = re.compile(r'\b\d+(\.\d+)?\s*(ml|g|gr|oz|fl|l|kg|mg|pcs?)\b', re.I)
STOP = {'the', 'and', 'for', 'with', 'of', 'de', 'la', 'le', 'du', 'new', 'set',
        'kit', 'skin', 'face', 'ml', 'g', 'in', 'to', 'a', 'x'}


def key(b, n):
    s = asc(n)
    for w in re.split(r'[^a-z0-9]+', asc(b)):
        if len(w) > 2:
            s = s.replace(w, ' ')
    s = SIZE.sub(' ', s)
    return (re.sub(r'[^a-z0-9]', '', asc(b)) + '||' +
            ' '.join(sorted(w for w in re.sub(r'[^a-z0-9 ]', ' ', s).split()
                            if w not in STOP)))


# ------------------------------------------------ 5a. tidy the brand names
# The vendor field arrives in several states and the same maker appears under
# several spellings, which would count one Lebanese brand as four:
#
#   Helwe / Helwehnaturalcosmetics / Helwe Natural Skincare / helwe
#   Sama Soaps / SamaSoaps
#   AloeLab / The AloeLab
#   Senteurs d'Orient / Senteurs d'Orient   (two different apostrophes)
#   HOUSE OF SOAP &#8211; Lebanon &#8211;   (undecoded HTML entities)
#   Cosmaline Online Shop                   (a shop name, not a brand)
import html as _html

TRAIL = re.compile(
    r'\s*(natural\s+)?(cosmetics?|skincare|skin care|beauty|care products?|'
    r'products?|online shop|shop|store|lebanon|lb|official|industry|'
    r'beauty care products?|products experts in lebanon)\s*$', re.I)


def tidy_brand(b):
    t = _html.unescape(str(b))
    t = re.sub(r'[‘’ʼ]', "'", t)          # curly apostrophes
    t = re.sub(r'[–—]', '-', t)                 # dashes
    t = re.sub(r'\s*[-|,].*$', '', t)                     # anything after a dash
    t = re.sub(r'\s+', ' ', t).strip(' -&')
    for _ in range(3):
        t = TRAIL.sub('', t).strip()
    return t or str(b)[:30]


both['brand'] = both['brand'].map(tidy_brand)

# merge spellings that are the same brand once accents and spacing are removed
def ckey(b):
    # a leading "the" has to go, or "The AloeLab" and "AloeLab" stay apart
    return re.sub(r'^the', '', re.sub(r'[^a-z0-9]', '', asc(b)))


canon = {}
for b in sorted(both['brand'].unique(), key=len):
    k = ckey(b)
    hit = next((canon[c] for c in canon
                if c and (k.startswith(c) or c.startswith(k)) and
                min(len(k), len(c)) >= 4), None)
    canon[k] = hit or b
both['brand'] = [canon.get(ckey(b), b) for b in both['brand']]

both['_k'] = [key(b, n) for b, n in zip(both['brand'], both['name'])]
before = len(both)
both = both.sort_values('ingredients', key=lambda s: s.str.len(), ascending=False)
both = both.drop_duplicates('_k').drop(columns=['_k']).reset_index(drop=True)
both['product_id'] = ['LO' + str(i + 1).zfill(5) for i in range(len(both))]
print(f'duplicates removed              {before - len(both):,}\n')

both.to_csv(OUT_C, index=False)

W = {'product_id': 10, 'brand': 24, 'name': 46, 'product_type': 18, 'country': 12,
     'ingredients': 55, 'product_url': 40, 'domain': 24, 'origin_evidence': 40,
     'harvest_method': 18, 'skin_type_url': 36, 'benefits': 34}
with pd.ExcelWriter(OUT_X, engine='openpyxl') as w:
    both.to_excel(w, sheet_name='Lebanese origin', index=False)
    ws = w.sheets['Lebanese origin']
    for c in ws[1]:
        v = str(c.value)
        c.fill = PatternFill('solid', fgColor='6f9c78' if v in EXTRA else '584A7A')
        c.font = Font(bold=True, color='FFFFFF', size=10)
        c.alignment = Alignment(vertical='center', horizontal='left')
        ws.column_dimensions[get_column_letter(c.column)].width = W.get(v, 16)
    ws.freeze_panes = 'D2'
    ws.auto_filter.ref = ws.dimensions

M = len(both)
print('=' * 66)
print('  LEBANESE ORIGIN DATASET')
print('=' * 66)
print(f'  harvested                 {start:,}')
print(f'  after removing shops, food, electronics and non-skincare   {M:,}')
print(f'  brands                    {both["brand"].nunique()}')
print(f'  with ingredients          {int((both["ingredients"] != "").sum()):,}'
      f'  ({100*(both["ingredients"] != "").mean():.0f}%)')
print(f'  with a skin type          {int((both["skin_type"] != "").sum()):,}'
      f'  ({100*(both["skin_type"] != "").mean():.0f}%)')
print(f'  with a price              {int((both["price_usd"] != "").sum()):,}'
      f'  ({100*(both["price_usd"] != "").mean():.0f}%)')
print()
print('  BRANDS')
for b, n in both['brand'].value_counts().items():
    print(f'    {b[:34]:34s}{n:5,}')
print()
print('  CATEGORIES')
for b, n in both['product_type'].value_counts().items():
    print(f'    {b[:24]:24s}{n:5,}')
print(f'\n  written to {OUT_X} and {OUT_C}')
print(f'  everything removed is in {DROPPED}, with the reason, so it is checkable')
