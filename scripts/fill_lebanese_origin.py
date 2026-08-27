"""
Fill the lebanese origin dataset from the brands' own pages

Why this one should do well
  Every row already carries product_url, and that URL is the BRAND'S OWN
  product page, because the catalogue was harvested from the brand's website.

Usage:
    py fill_lebanese_origin.py                 free, the brands' own pages
    py fill_lebanese_origin.py --test 100      measure first, write nothing
"""
import os
import re
import sys
import csv
import html as htmllib
import threading
import requests
import pandas as pd
from concurrent.futures import ThreadPoolExecutor, as_completed

DATA = 'LEBANESE_ORIGIN.csv'
PROGRESS = 'lebanese_origin_fill_progress.csv'
PFIELDS = ['product_id', 'ingredients', 'skin_type', 'sensitivity',
           'skin_type_quote', 'skin_type_rule', 'product_summary']
WORKERS = 8
TIMEOUT = 20
TEST = 0
if '--test' in sys.argv:
    i = sys.argv.index('--test')
    TEST = int(sys.argv[i + 1]) if len(sys.argv) > i + 1 else 100

HEAD = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 '
                      '(KHTML, like Gecko) Chrome/124.0 Safari/537.36',
        'Accept-Language': 'en-US,en;q=0.9,fr;q=0.8,ar;q=0.7'}

TAG = re.compile(r'<[^>]+>')
HEADING = re.compile(r'(ingredients?|composition|inci|ingr[ée]dients?|'
                     r'contents|formula|المكونات|تركيبة)', re.I)
LEADIN = re.compile(r'\b(?:full\s+|main\s+|key\s+|active\s+|other\s+)?'
                    r'(?:ingredients?|inci|composition|contents)\b\s*[:\-–]?\s*', re.I)
INSTRUCTION = re.compile(r'\b(apply|rinse|massage|leave on|caution|patch test|'
                         r'for external use|avoid contact|shake well)\b', re.I)
JUNK_PIECE = re.compile(r'(app store|google play|add to cart|checkout|©|http|'
                        r'shipping|newsletter|\blogin\b)', re.I)
KNOWN = re.compile(
    r'\b(aqua|water|eau|glycerin|glycerine|sodium|potassium|calcium|zinc|'
    r'titanium|cetearyl|cetyl|stearyl|stearic|lauryl|laureth|dimethicone|'
    r'phenoxyethanol|tocopherol|parfum|fragrance|limonene|linalool|citral|'
    r'geraniol|coumarin|xanthan|carbomer|panthenol|niacinamide|retinol|'
    r'ascorbic|hyaluron|salicylic|glycolic|lactic|citric|benzoic|sorbic|'
    r'caprylic|capric|triglyceride|polysorbate|peg-|butylene|propylene|'
    r'pentylene|hexanediol|urea|allantoin|bisabolol|squalane|shea|olea|cocos|'
    r'butyrospermum|simmondsia|prunus|helianthus|aloe|centella|camellia|'
    r'chamomilla|lavandula|rosa|vitis|citrus|mineral oil|paraffinum|'
    r'petrolatum|lanolin|beeswax|cera alba|laurus|niacin|ethylhexyl|'
    r'homosalate|avobenzone|octocrylene|edta|paraben|benzyl|caffeine|'
    r'adenosine|arbutin|kojic|azelaic|ceramide|collagen|keratin|'
    r'hydroxyethyl|hydroxypropyl|polyquaternium|acrylates|copolymer|ethanol|'
    r'isopropyl|isododecane|dicaprylyl|coco-|decyl|myristate|palmitate|'
    r'oleate|linoleate|behenate|honey|mel\b|propolis)', re.I)
CHEM_END = re.compile(
    r'(ate|ite|ide|ine|one|ol|oic acid|ic acid|yl|ene|ose|an|um|ium|extract|'
    r'oil|butter|wax|water|acid|glycol|glycerin|alcohol|gum|powder|seed|leaf|'
    r'root|fruit|flower|juice|ferment|filtrate|protein|peptide|ceramide|'
    r'vitamin|silica|mica|dioxide|oxide)$', re.I)
VERB = re.compile(
    r'\b(is|are|was|were|be|been|has|have|helps?|works?|leaves?|makes?|gives?|'
    r'provides?|delivers?|targets?|boosts?|reduces?|improves?|designed|'
    r'formulated|created|apply|rinse|use|discover|our|your|we |this |that |'
    r'which |thanks to)\b', re.I)
BENEFIT = re.compile(
    r'\s*(hydrating|moisturi[sz]ing|brightening|anti[- ]?aging|anti[- ]?ageing|'
    r'soothing|calming|firming|smoothing|nourishing|purifying|clarifying|'
    r'exfoliating|cleansing|mattifying|plumping|lifting|barrier repair|'
    r'oil control|radiance|glow)\s*', re.I)
GAP = 2

TYPE_W = {'dry': 'Dry', 'oily': 'Oily', 'oil': 'Oily', 'combination': 'Combination',
          'combo': 'Combination', 'normal': 'Normal', 'acne': 'Oily'}
UNIV = re.compile(r'\b(all skin types?|every skin type|any skin type|'
                  r'tous types de peaux?)\b', re.I)
CLAIM = [
    (1, 'declared field',
     re.compile(r'skin[\s_-]*type[^a-z0-9]{0,4}\s*[:\-]\s*([A-Za-z ,/&+-]{3,60})', re.I)),
    (2, 'suitability sentence',
     re.compile(r'\b(?:suitable|suited|recommended|formulated|designed|ideal|perfect|'
                r'great|made|created|convient)\s+(?:for|to|aux?)\s+'
                r'([a-z ,/&+-]{3,70}?)\s*(?:skin|peaux?)\b', re.I)),
    (2, 'reversed phrase',
     re.compile(r'\b([a-z ,/&+-]{3,50}?)\s*skin\s*types?\b', re.I)),
    (2, 'soothing wording',
     re.compile(r'\b(?:soothe?s?|soothing|calm(?:s|ing)?|gentle|non[- ]irritating|'
                r'fragrance[- ]free|hypoallergenic)\b[^.]{0,60}?'
                r'\b(irritated|reactive|delicate|sensitive)\b', re.I)),
]


def clean(t):
    return re.sub(r'\s+', ' ', htmllib.unescape(TAG.sub(' ', str(t)))).strip(' :.-•|')


def is_ing(p):
    q = p.strip(' .;:•|*-()[]')
    if not q or len(q) > 70 or JUNK_PIECE.search(q) or VERB.search(q):
        return False
    w = q.split()
    if len(w) > 7 or q.count('.') > 1:
        return False
    if KNOWN.search(q):
        return True
    last = re.sub(r'[^A-Za-z]', '', w[-1]) if w else ''
    if len(last) >= 3 and CHEM_END.search(last):
        return True
    if re.match(r'^(ci\s*\d{4,6}|[a-z]{2,}-\d+)$', q, re.I):
        return True
    return len(w) <= 5 and bool(re.fullmatch(r"[^\W\d_][\w \-/()'.&%+]*", q, re.UNICODE))


def find_inci(text):
    t = clean(text)
    t = re.sub(r'\b(full\s+)?(inci\s*(formula|list)?|active ingredients?|'
               r'inactive ingredients?|ingredients?|composition)\s*[:\-–]\s*',
               ', ', t, flags=re.I)
    pieces = [p for p in re.split(r'[,;]|\.\s+|\.$|\s[-–—]\s|\s•\s', t)]
    flags = [is_ing(p) for p in pieces]
    spans, start, gap, last = [], -1, 0, -1
    for i, ok in enumerate(flags + [False] * (GAP + 1)):
        if ok:
            if start < 0:
                start = i
            gap, last = 0, i
        elif start >= 0:
            gap += 1
            if gap > GAP:
                spans.append((start, last, sum(flags[start:last + 1])))
                start, gap = -1, 0
    if not spans:
        return ''
    a, b, good = max(spans, key=lambda s: s[2])
    if good < 5:
        return ''
    kept = ', '.join(p.strip(' .;:•|*-') for p in pieces[a:b + 1] if p.strip(' .;:•|*-'))
    chem = sum(1 for p in pieces[a:b + 1]
               if KNOWN.search(p) or (p.split() and
                                      len(re.sub(r'[^A-Za-z]', '', p.split()[-1])) >= 3 and
                                      CHEM_END.search(re.sub(r'[^A-Za-z]', '', p.split()[-1]))))
    if chem < 3:
        return ''
    if sum(1 for p in pieces[a:b + 1] if BENEFIT.fullmatch(p.strip(' .;:•|*-'))) >= 3:
        return ''
    if INSTRUCTION.search(kept[:120]):
        return ''
    return kept[:4000]


def n_types(p):
    t = str(p).lower()
    return sum(bool(re.search(r'\b' + w, t))
               for w in ('dry', 'oil', 'combination', 'normal', 'sensitiv'))


def parse_types(t):
    t = str(t).lower()
    f = {v for k, v in TYPE_W.items() if re.search(r'\b' + k, t)}
    sn = 'Sensitive' if re.search(r'\b(sensitiv|irritated|reactive|delicate)', t) else ''
    if ('Dry' in f and 'Oily' in f) or 'Combination' in f:
        return 'Combination', sn
    for lab in ('Oily', 'Dry', 'Normal'):
        if lab in f:
            return lab, sn
    return '', sn


def sentence(flat, a, b):
    i = max(flat.rfind('.', 0, a), flat.rfind('|', 0, a), 0)
    j = min([x for x in (flat.find('.', b), flat.find('|', b)) if x != -1] or [len(flat)])
    return flat[i:j + 1].strip(' .|').strip()


def find_claim(page):
    flat = re.sub(r'\s+', ' ', clean(
        re.sub(r'<script[\s\S]*?</script>|<style[\s\S]*?</style>', ' ', page)))
    for strength, rule, pat in CLAIM:
        for m in pat.finditer(flat):
            phrase = m.group(1)
            q = sentence(flat, m.start(), m.end())
            if len(q) < 12 or len(q) > 300 or JUNK_PIECE.search(q):
                continue
            if n_types(phrase) >= 4 or UNIV.search(phrase) or UNIV.search(q):
                return 'All', '', rule + ' (all skin types)', q[:220]
            st, sn = parse_types(phrase)
            if st or sn:
                return st, sn, rule, q[:220]
    return '', '', '', ''


S = requests.Session()
S.headers.update(HEAD)
lock = threading.Lock()

df = pd.read_csv(DATA, dtype=str, low_memory=False).fillna('')
N = len(df)
for c in ('ingredients_raw', 'ingredient_source', 'ingredient_url'):
    if c not in df.columns:
        df[c] = ''

already = set()
if os.path.exists(PROGRESS) and not TEST:
    try:
        p = pd.read_csv(PROGRESS, dtype=str).fillna('')
        byid = {v: k for k, v in df['product_id'].items()}
        n = 0
        for _, r in p.iterrows():
            already.add(r['product_id'])
            i = byid.get(r['product_id'])
            if i is None:
                continue
            for c in PFIELDS[1:]:
                if r.get(c, '') and not df.at[i, c]:
                    df.at[i, c] = r[c]
                    n += 1
        print(f'resuming: {len(already):,} already read\n')
    except Exception as e:
        print(f'(progress unreadable: {e})\n')

need = df[(df['product_url'].str.startswith('http')) &
          ((df['ingredients'] == '') | (df['skin_type'] == '')) &
          (~df['product_id'].isin(already))]
if TEST:
    need = need.sample(min(TEST, len(need)), random_state=5)
    print(f'TEST MODE: {len(need)} products, nothing will be saved\n')
else:
    print(f'{N:,} Lebanese origin products')
    print(f'{len(need):,} need ingredients or a skin type and have a brand URL')
    print('reading the brands\' own pages. no search credits are used.\n')

found_i, found_s, pending, done = {}, {}, [], [0]


def flush(force=False):
    global pending
    if not pending or (len(pending) < 50 and not force):
        return
    new = not os.path.exists(PROGRESS)
    with open(PROGRESS, 'a', newline='', encoding='utf-8') as fh:
        w = csv.DictWriter(fh, fieldnames=PFIELDS)
        if new:
            w.writeheader()
        w.writerows(pending)
    pending = []


def work(i):
    url = str(df.at[i, 'product_url'])
    try:
        r = S.get(url, timeout=TIMEOUT)
        if r.status_code != 200:
            return
        page = r.text[:400000]
    except Exception:
        return
    rec = {'product_id': df.at[i, 'product_id']}
    if not df.at[i, 'ingredients']:
        ing = find_inci(page)
        if ing:
            rec['ingredients'] = ing
            with lock:
                found_i[i] = ing
    if not df.at[i, 'skin_type']:
        st, sn, rule, quote = find_claim(page)
        if st or sn:
            rec.update(skin_type=st or 'All', sensitivity=sn or 'Resistant',
                       skin_type_rule=rule, skin_type_quote=quote)
            with lock:
                found_s[i] = (st or 'All', sn or 'Resistant', rule, quote)
    if not df.at[i, 'product_summary']:
        m = re.search(r'<meta[^>]+name=["\']description["\'][^>]+content=["\']([^"\']{40,400})',
                      page, re.I)
        if m:
            rec['product_summary'] = clean(m.group(1))[:400]
    with lock:
        done[0] += 1
        if len(rec) > 1 and not TEST:
            pending.append({k: rec.get(k, '') for k in PFIELDS})
            flush()
        if done[0] % 100 == 0:
            print(f'   {done[0]:,}/{len(need):,}   ingredients {len(found_i):,}'
                  f'   skin type {len(found_s):,}', flush=True)


with ThreadPoolExecutor(max_workers=WORKERS) as ex:
    list(ex.map(work, need.index))
if not TEST:
    flush(force=True)

print(f'\n   read {len(need):,} brand pages')
print(f'   ingredients found  {len(found_i):,}  ({100*len(found_i)/max(len(need),1):.0f}%)')
print(f'   skin types found   {len(found_s):,}  ({100*len(found_s)/max(len(need),1):.0f}%)')
print('\n   EXAMPLES')
for i in list(found_i)[:4]:
    print(f'      {df.at[i,"brand"][:16]:16s} {df.at[i,"name"][:34]:34s}')
    print(f'         {found_i[i][:88]}')
for i in list(found_s)[:3]:
    st, sn, rule, q = found_s[i]
    print(f'      {df.at[i,"brand"][:16]:16s} {df.at[i,"name"][:30]:30s} -> {st}/{sn}')
    print(f'         "{q[:80]}"')

if TEST:
    print('\n' + '=' * 60)
    print('  TEST ONLY, nothing written.')
    n_i = int((df['ingredients'] == '').sum())
    n_s = int((df['skin_type'] == '').sum())
    print(f'  a full run would add about {int(n_i*len(found_i)/max(len(need),1)):,} '
          f'ingredient lists and {int(n_s*len(found_s)/max(len(need),1)):,} skin types')
    raise SystemExit

for i, ing in found_i.items():
    df.at[i, 'ingredients'] = ing
    df.at[i, 'ingredients_raw'] = ing
    df.at[i, 'ingredient_source'] = df.at[i, 'domain']
    df.at[i, 'ingredient_url'] = df.at[i, 'product_url']
    df.at[i, 'ingredient_count'] = str(len([x for x in ing.split(',') if x.strip()]))
for i, (st, sn, rule, q) in found_s.items():
    df.at[i, 'skin_type'] = st
    df.at[i, 'sensitivity'] = sn
    df.at[i, 'skin_type_rule'] = rule
    df.at[i, 'skin_type_quote'] = q
    df.at[i, 'skin_type_tier'] = '1'
    df.at[i, 'skin_type_authority'] = 'manufacturer'
    df.at[i, 'skin_type_source'] = df.at[i, 'domain']
    df.at[i, 'skin_type_url'] = df.at[i, 'product_url']
    df.at[i, 'skin_type_status'] = 'found'
df.loc[df['skin_type'] == '', 'skin_type_status'] = 'searched, not stated'

df.to_csv(DATA, index=False)

try:
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    EXTRA = {'brand_country', 'market', 'price_usd', 'product_url', 'domain',
             'harvest_method', 'origin_evidence'}
    with pd.ExcelWriter('LEBANESE_ORIGIN.xlsx', engine='openpyxl') as w:
        df.to_excel(w, sheet_name='Lebanese origin', index=False)
        ws = w.sheets['Lebanese origin']
        for c in ws[1]:
            v = str(c.value)
            c.fill = PatternFill('solid', fgColor='6f9c78' if v in EXTRA else '584A7A')
            c.font = Font(bold=True, color='FFFFFF', size=10)
            c.alignment = Alignment(vertical='center', horizontal='left')
            ws.column_dimensions[get_column_letter(c.column)].width = 18
        ws.freeze_panes = 'D2'
        ws.auto_filter.ref = ws.dimensions
except PermissionError:
    print('(workbook open in Excel, csv written)')

print('\n' + '=' * 62)
print('  LEBANESE ORIGIN, AFTER READING THE BRAND PAGES')
print('=' * 62)
for c in ('brand', 'name', 'product_type', 'ingredients', 'skin_type',
          'sensitivity', 'price_usd', 'product_summary', 'product_url'):
    n = int((df[c] != '').sum())
    print(f'   {c:18s}{n:6,}  ({100*n/N:5.1f}%)')
print(f'\n   {N:,} products, {df["brand"].nunique()} brands')
